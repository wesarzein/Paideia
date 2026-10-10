"""Read the privately inspected institutional template, without evaluating Excel code.

Block ordinals are not calendar months. Derived cells are reference information only.
The descriptor deliberately separates source structure from academic configuration.
"""

import re
import unicodedata
from decimal import ROUND_HALF_UP, Decimal
from io import BytesIO

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from app.services.evidence_types import normalize_evidence_type


def key(value):
    text = unicodedata.normalize("NFKD", str(value or ""))
    return re.sub(r"[^a-z0-9]", "", text.encode("ascii", "ignore").decode().lower())


def read_institutional(content: bytes) -> dict | None:
    workbook = load_workbook(BytesIO(content), data_only=False, keep_links=False)
    cached = load_workbook(BytesIO(content), data_only=True, keep_links=False)
    matches = []
    for sheet in workbook:
        if sheet.max_row > 10000 or sheet.max_column > 200:
            raise ValueError("Hoja demasiado grande para la vista previa")
        for cells in sheet.iter_rows(max_row=min(30, sheet.max_row)):
            codes = [c for c in cells if key(c.value) in {"codigo", "studentcode"}]
            names = [c for c in cells if key(c.value) in {"nombreyapellidos", "apellidosynombres"}]
            blocks = [c for c in cells if re.fullmatch(r"[12]mes", key(c.value))]
            if codes and names and len(blocks) == 2:
                matches.append((sheet, codes[0], names[0], blocks))
    if not matches:
        return None  # Preserve the existing flat CSV/XLSX import.
    if len(matches) != 1:
        raise ValueError(
            "Varias hojas institucionales: separa o selecciona la hoja antes de importar"
        )
    sheet, code_cell, name_cell, parents = matches[0]
    descriptor, inputs, errors, derived = [], [], [], []
    for parent in parents:
        merged = next((m for m in sheet.merged_cells.ranges if parent.coordinate in m), None)
        if merged is None:
            raise ValueError("El bloque mensual no tiene un encabezado combinado reconocible")
        criteria = []
        for column in range(merged.min_col, merged.max_col + 1):
            header = sheet.cell(parent.row + 1, column).value
            try:
                criterion = normalize_evidence_type(str(header or ""))
            except ValueError:
                continue  # P.S / PROMEDIO MENSUAL are not input criteria.
            criteria.append({"column": column, "name": criterion, "header": str(header)})
        if len(criteria) != 5 or len({c["name"] for c in criteria}) != 5:
            raise ValueError("Bloque con criterios desconocidos; requiere mapeo de plantilla")
        descriptor.append(
            {
                "block": key(parent.value)[0],
                "label": parent.value,
                "criteria": criteria,
                "summary_column": merged.max_col,
            }
        )
    blank_inputs = 0
    students = []
    unsupported_formulas = []
    for source_row in range(code_cell.row + 2, sheet.max_row + 1):
        code = sheet.cell(source_row, code_cell.column).value
        name = sheet.cell(source_row, name_cell.column).value
        if not code and not name:
            continue
        # Ignore repeated auxiliary headers, not arbitrary rows containing invalid marks.
        if key(code) in {"codigo", "studentcode"}:
            continue
        expected = {}
        summaries = []
        for block in descriptor:
            letters = [f"{get_column_letter(c['column'])}{source_row}" for c in block["criteria"]]
            summary = f"{get_column_letter(block['summary_column'])}{source_row}"
            summaries.append(summary)
            total = f"SUM({letters[0]}:{letters[2]})"
            if block["block"] == "2":
                total = f"({total})"
            expected[summary] = f"=ROUND(({total}+{letters[3]}*1+{letters[4]}*1)/5,1)"
        finals = [c for c in sheet[code_cell.row] if key(c.value) == "pf"]
        literals = [c for c in sheet[code_cell.row] if key(c.value) == "promedioliteral"]
        if len(finals) == len(literals) == 1:
            final_cell = f"{finals[0].column_letter}{source_row}"
            literal_cell = f"{literals[0].column_letter}{source_row}"
            expected[final_cell] = f"=ROUND(AVERAGE({summaries[0]},{summaries[1]}),1)"
            expected[literal_cell] = (
                f'=IF({final_cell}>16.4,"AD",IF({final_cell}>13.4,"A",'
                f'IF({final_cell}>10.4,"B","C")))'
            )
        verified = len(expected) == 4 and all(
            re.sub(r"\s", "", str(sheet[position].value)).upper() == formula
            for position, formula in expected.items()
        )
        students.append(
            {"source_row": source_row, "student_code": str(code or ""), "name": str(name or "")}
        )
        if not verified:
            unsupported_formulas.append(source_row)
        for block in descriptor:
            for criterion in block["criteria"]:
                cell = sheet.cell(source_row, criterion["column"])
                if cell.value is None or str(cell.value).strip() == "":
                    blank_inputs += 1
                    continue
                value = cell.value
                row = {
                    "student_code": str(code or "").strip(),
                    "source_name": str(name or ""),
                    "score": value,
                    "evaluation_name": "",
                    "evaluation_type": criterion["name"],
                    "assessment_date": "",
                    "qualitative_note": "",
                    "source_block": block["block"],
                    "source_cell": cell.coordinate,
                    "source_row": source_row,
                    "source_sheet": sheet.title,
                    "template": "institutional_v1",
                }
                row["reference_verified"] = verified
                inputs.append(row)
                try:
                    if (
                        cell.data_type == "f"
                        or isinstance(value, bool)
                        or not 0 <= float(value) <= 20
                    ):
                        raise ValueError
                except (TypeError, ValueError):
                    errors.append(
                        {
                            "row": len(inputs),
                            "error": "Entrada inválida; corrige la nota (0–20), no una fórmula",
                        }
                    )
        for cells in sheet.iter_rows(min_row=source_row, max_row=source_row):
            for cell in cells:
                if cell.data_type == "f" and all(
                    cell.column not in [c["column"] for c in b["criteria"]] for b in descriptor
                ):
                    derived.append(
                        {
                            "cell": cell.coordinate,
                            "formula": cell.value,
                            "cached": cached[sheet.title][cell.coordinate].value,
                            "source_row": source_row,
                        }
                    )
    return {
        "columns": [
            "student_code",
            "source_name",
            "source_block",
            "evaluation_type",
            "score",
            "evaluation_name",
            "assessment_date",
            "qualitative_note",
        ],
        "rows": inputs,
        "total": len(inputs),
        "errors": errors,
        "institutional": {
            "template": "institutional_v1",
            "sheet": sheet.title,
            "blocks": descriptor,
            "derived": derived,
            "blank_inputs": blank_inputs,
            "students": students,
            "unsupported_formula_rows": unsupported_formulas,
            "warning": (
                "Vacíos omitidos; cachés y literales no son notas. "
                "Confirma calendario, contexto y correspondencias."
            ),
        },
    }


def reference_results(rows: list[dict]) -> list[dict]:
    """Inspected formula reference, never execute file expressions or persist aggregates.

    Only complete sets get a result. Excel SUM treats blanks as zero, but that result
    is deliberately not presented as an actual mark for an incomplete evaluation.
    """
    groups = {}
    for row in rows:
        if row.get("template") != "institutional_v1":
            continue
        group = groups.setdefault((row.get("source_sheet"), row.get("source_row")), {})
        if not row.get("reference_verified"):
            group["unsupported"] = True
        group.setdefault(str(row.get("source_block")), {})[row["evaluation_type"]] = row["score"]
    results = []
    for (_sheet, source_row), blocks in groups.items():
        monthly = {}
        for block in ("1", "2"):
            criteria = blocks.get(block, {})
            if len(criteria) != 5 or blocks.get("unsupported"):
                monthly[block] = None
                continue
            try:
                values = [Decimal(str(value)) for value in criteria.values()]
                if any(not value.is_finite() or not 0 <= value <= 20 for value in values):
                    raise ValueError
                monthly[block] = (sum(values) / 5).quantize(Decimal("0.1"), ROUND_HALF_UP)
            except (ValueError, ArithmeticError):
                monthly[block] = None
        final = None
        if all(value is not None for value in monthly.values()):
            final = ((monthly["1"] + monthly["2"]) / 2).quantize(Decimal("0.1"), ROUND_HALF_UP)
        literal = (
            None
            if final is None
            else next(
                (
                    label
                    for threshold, label in [
                        (Decimal("16.4"), "AD"),
                        (Decimal("13.4"), "A"),
                        (Decimal("10.4"), "B"),
                    ]
                    if final > threshold
                ),
                "C",
            )
        )
        results.append(
            {
                "source_row": source_row,
                "monthly": {b: float(v) if v is not None else None for b, v in monthly.items()},
                "final": float(final) if final is not None else None,
                "literal": literal,
                "reference_only": True,
            }
        )
    return results

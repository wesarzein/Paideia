"""Authorization across intersecting courses/sections and linked family accounts.

The guarded fixture uses SQLite or a unique schema in the authorized synthetic PG.
No production session, schema public or existing laboratory records are modified.
"""

from datetime import date
from io import BytesIO
from types import SimpleNamespace
from uuid import uuid4

import pytest
from fastapi import FastAPI
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from test_roster_persistence import sessions as _sessions

from app.api.dependencies import get_current_user
from app.core.database import get_db
from app.endpoints import (
    academic,
    attendance,
    dashboard,
    enrollments,
    follow_up,
    grades,
    reports,
    students,
)
from app.models.academic import (
    AcademicPeriod,
    AttendanceRecord,
    Course,
    CourseAssignment,
    Enrollment,
    Grade,
    GradeRecord,
    Section,
)
from app.models.follow_up import FollowUp
from app.models.student import Student
from app.models.user import Role, User

sessions = _sessions


@pytest.fixture
def scenario(sessions):
    with sessions() as db:
        actors = {}
        for code in ("admin", "teacher", "coordinator", "director", "student", "parent"):
            role = Role(code=code, name=code)
            user = User(
                id=uuid4(),
                role=role,
                email=f"{code}@synthetic.invalid",
                full_name=code,
                password_hash="synthetic",
            )
            db.add_all([role, user])
            actors[code] = SimpleNamespace(id=user.id, role=SimpleNamespace(code=code))
        db.flush()
        grade = Grade(id=uuid4(), name="1°", level="Secundaria")
        period = AcademicPeriod(
            id=uuid4(),
            name="Año escolar 2026",
            starts_on=date(2026, 1, 1),
            ends_on=date(2026, 12, 31),
        )
        sections = [Section(id=uuid4(), grade_id=grade.id, name=name) for name in ("A", "B", "C")]
        courses = [Course(id=uuid4(), name=name, code=name) for name in ("Scope X", "Scope Y")]
        pupils = [
            Student(
                id=uuid4(),
                student_code=f"SYN-{index}",
                first_name=f"Alumno {index}",
                last_name="Sintético",
                grade_id=grade.id,
                section_id=section.id,
            )
            for index, section in enumerate(sections)
        ]
        pupils[0].student_user_id = actors["student"].id
        pupils[0].parent_user_id = pupils[1].parent_user_id = actors["parent"].id
        db.add_all([grade, period, *sections, *courses, *pupils])
        db.flush()
        db.add_all(
            [
                CourseAssignment(
                    course_id=courses[index].id,
                    grade_id=grade.id,
                    section_id=sections[index].id,
                    teacher_id=actors["teacher"].id,
                )
                for index in (0, 1)
            ]
        )
        enrollment_ids = []
        for index, (pupil, course, score, status) in enumerate(
            [
                (pupils[0], courses[0], 10, "PRESENT"),
                (pupils[0], courses[1], 20, "ABSENT"),
                (pupils[1], courses[1], 12, "PRESENT"),
                (pupils[1], courses[0], 20, "ABSENT"),
                (pupils[2], courses[0], 20, "ABSENT"),
            ]
        ):
            db.add(
                GradeRecord(
                    student_id=pupil.id,
                    course_id=course.id,
                    period_id=period.id,
                    evaluation_name=f"Prueba {index}",
                    evaluation_type="Cuaderno",
                    assessment_date=date(2026, 4, 1),
                    score=score,
                )
            )
            db.add(
                AttendanceRecord(
                    student_id=pupil.id,
                    course_id=course.id,
                    period_id=period.id,
                    attendance_date=date(2026, 4, 1),
                    status=status,
                )
            )
            enrollment = Enrollment(
                id=uuid4(),
                student_id=pupil.id,
                course_id=course.id,
                period_id=period.id,
                status="ACTIVE",
            )
            enrollment_ids.append(str(enrollment.id))
            db.add(enrollment)
        observations = [
            FollowUp(
                id=uuid4(),
                student_id=pupil.id,
                action="Observación sintética",
                recorded_by_id=actors["admin"].id,
            )
            for pupil in pupils
        ]
        db.add_all(observations)
        ids = {
            "students": [str(p.id) for p in pupils],
            "courses": [str(c.id) for c in courses],
            "sections": [str(s.id) for s in sections],
            "grade": str(grade.id),
            "period": str(period.id),
            "enrollments": enrollment_ids,
            "followups": [str(f.id) for f in observations],
        }
        db.commit()

    application = FastAPI()
    for module, prefix in [
        (academic, "/academic"),
        (attendance, "/attendance"),
        (dashboard, "/dashboard"),
        (enrollments, "/enrollments"),
        (follow_up, "/follow-up"),
        (grades, "/grades"),
        (reports, "/reports"),
        (students, "/students"),
    ]:
        application.include_router(module.router, prefix=prefix)
    current = {"user": actors["teacher"]}

    def database():
        with sessions() as db:
            yield db

    application.dependency_overrides[get_db] = database
    application.dependency_overrides[get_current_user] = lambda: current["user"]
    with TestClient(application) as client:
        yield client, ids, actors, current


def test_teacher_records_and_aggregates_are_correlated_to_course_and_section(scenario):
    client, ids, _actors, _current = scenario
    for route in ("/grades", "/attendance", "/enrollments"):
        response = client.get(route)
        assert response.status_code == 200
        assert {(r["student_id"], r["course_id"]) for r in response.json()} == {
            (ids["students"][0], ids["courses"][0]),
            (ids["students"][1], ids["courses"][1]),
        }
    assert client.get(f"/grades/summary/{ids['students'][0]}").json()["average"] == 10
    for route in ("/dashboard/summary", "/reports/summary"):
        result = client.get(route)
        assert result.status_code == 200
        assert result.json()["total_students"] == 2
        assert result.json()["average_score"] == 11
        assert result.json()["attendance_rate"] == 100
    assert client.get(f"/grades/summary/{ids['students'][2]}").status_code == 403
    detail = client.get(f"/reports/students/{ids['students'][0]}").json()
    assert {r["course"] for r in detail["grades"]} == {"Scope X"}
    assert {r["course"] for r in detail["attendance"]} == {"Scope X"}
    assert client.get(f"/reports/students/{ids['students'][2]}").status_code == 403


def test_teacher_export_and_filters_never_include_an_unassigned_course_section(scenario):
    client, ids, _actors, _current = scenario
    filters = {"grade_id": ids["grade"], "course_id": ids["courses"][0]}
    rows = client.get("/reports/students", params=filters).json()
    assert [r["student_id"] for r in rows] == [ids["students"][0]]
    result = client.get("/reports/export.xlsx", params=filters)
    assert result.status_code == 200
    sheet = load_workbook(BytesIO(result.content)).active
    assert sheet.max_row == 2 and "Alumno 0" in str(sheet.cell(2, 3).value)
    assert client.get("/reports/students", params={"section_id": ids["sections"][2]}).json() == []
    assert client.get("/grades", params={"student_id": ids["students"][2]}).json() == []


@pytest.mark.parametrize("role", ["admin", "coordinator"])
def test_institutional_read_roles_can_consult_and_export_reports(scenario, role):
    client, ids, actors, current = scenario
    current["user"] = actors[role]
    assert client.get("/reports/summary").status_code == 200
    assert len(client.get("/reports/students").json()) == 3
    assert client.get("/reports/export.xlsx", params={"grade_id": ids["grade"]}).status_code == 200
    assert len(client.get("/grades").json()) == 5


@pytest.mark.parametrize("role", ["director", "student", "parent"])
def test_retired_roles_cannot_read_or_write_academic_records(scenario, role):
    client, ids, actors, current = scenario
    current["user"] = actors[role]
    for path in (
        "/grades", "/attendance", "/students", "/academic/catalog",
        "/reports/summary", "/reports/students", "/enrollments", "/follow-up",
        f"/students/{ids['students'][0]}",
        f"/grades/summary/{ids['students'][0]}",
    ):
        assert client.get(path).status_code == 403
    payload = {
        "student_id": ids["students"][0],
        "course_id": ids["courses"][0],
        "period_id": ids["period"],
        "score": 12,
        "assessment_date": "2026-04-01",
    }
    assert client.post("/grades", json=payload).status_code == 403


def test_teacher_cannot_enroll_or_modify_records_outside_assignment(scenario):
    client, ids, _actors, _current = scenario
    payload = {
        "student_id": ids["students"][0],
        "course_id": ids["courses"][1],
        "period_id": ids["period"],
        "score": 12,
        "assessment_date": "2026-04-01",
    }
    assert client.post("/grades", json=payload).status_code == 403
    assert client.post("/enrollments", json=payload).status_code == 403
    assert client.get(f"/enrollments/{ids['enrollments'][1]}").status_code == 403
    assert client.get(f"/enrollments/{ids['enrollments'][0]}").status_code == 200
    assert client.get("/follow-up").status_code == 403
    assert (
        client.patch(
            f"/follow-up/{ids['followups'][2]}", json={"action": "Cambio ajeno"}
        ).status_code
        == 403
    )
    assert client.delete(f"/follow-up/{ids['followups'][2]}").status_code == 403


def test_historical_inconsistent_grade_section_does_not_grant_teacher_access(scenario, sessions):
    client, ids, _actors, _current = scenario
    # Legacy rows can contain a grade inconsistent with their section FK.
    from uuid import UUID

    with sessions() as db:
        other_grade = Grade(id=uuid4(), name="2°", level="Secundaria")
        db.add(other_grade)
        db.flush()
        pupil = Student(
            id=uuid4(),
            student_code="SYN-INCONSISTENT",
            first_name="Prueba",
            last_name="Histórica",
            grade_id=other_grade.id,
            section_id=UUID(ids["sections"][0]),
        )
        db.add(pupil)
        db.flush()
        db.add(
            GradeRecord(
                student_id=pupil.id,
                course_id=UUID(ids["courses"][0]),
                period_id=UUID(ids["period"]),
                assessment_date=date(2026, 4, 1),
                evaluation_name="Legado inconsistente",
                score=20,
            )
        )
        db.commit()
    assert len(client.get("/grades").json()) == 2
    assert client.get("/dashboard/summary").json()["average_score"] == 11

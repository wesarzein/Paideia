"""Institutional XLSX integration, synthetic pupils, isolated PostgreSQL only."""

import os
from copy import deepcopy
from io import BytesIO
from uuid import uuid4

import pytest
from institutional_fixture import workbook_bytes
from openpyxl import load_workbook
from openpyxl.worksheet.cell_range import CellRange
from sqlalchemy import select
from test_grade_import_identity import client_for
from test_roster_persistence import roster
from test_roster_persistence import sessions as _sessions

from app.models.academic import (
    AcademicPeriod,
    Course,
    CourseAssessmentComponent,
    CourseAssignment,
    GradeRecord,
)
from app.models.student import Student
from app.services import school_roster
from app.services.institutional_excel import read_institutional, reference_results

sessions = _sessions
pytestmark = pytest.mark.skipif(
    not os.environ.get("PAIDEIA_TEST_POSTGRES_URL"), reason="Requires disposable PostgreSQL"
)


def setup(sessions):
    with sessions() as db:
        school_roster.synchronize_roster(
            db, roster(students=[f"SINTETICO, ALUMNO {i:03d}" for i in range(1, 17)])
        )
        db.commit()
        pupil = db.scalar(select(Student))
        return {
            "course_id": str(db.scalar(select(Course)).id),
            "period_id": str(db.scalar(select(AcademicPeriod)).id),
            "academic_context": {
                "grade_id": str(pupil.grade_id),
                "section_id": str(pupil.section_id),
                "school_year": 2026,
                "bimestre": 1,
                "confirmed": True,
                "blocks": {
                    "1": {
                        "evaluation_name": "Bloque abril confirmado",
                        "assessment_date": "2026-04-01",
                    },
                    "2": {
                        "evaluation_name": "Bloque mayo confirmado",
                        "assessment_date": "2026-04-29",
                    },
                },
            },
        }


def preview(client, content):
    response = client.post(
        "/imports/grades/preview",
        files={
            "file": (
                "institucional-sintetico.xlsx",
                content,
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
        },
    )
    assert response.status_code == 200, response.text
    return response.json()


def test_complete_template_160_inputs_persist_reimport_and_snapshots(sessions):
    payload = setup(sessions)
    with client_for(sessions) as client:
        result = preview(client, workbook_bytes("full"))
        assert result["total"] == 160 and len(result["institutional"]["derived"]) == 64
        assert {r["source_block"] for r in result["rows"]} == {"1", "2"}
        assert all(r["reference_verified"] for r in result["rows"])
        assert all(not r["assessment_date"] for r in result["rows"])
        payload["rows"] = result["rows"]
        review = client.post("/imports/grades/review", json=payload)
        assert review.status_code == 200, review.text
        assert review.json()["reference_results"][0]["final"] == 17
        assert review.json()["reference_results"][0]["literal"] == "AD"
        assert client.post("/imports/grades/commit", json=payload).json()["imported"] == 160
        before = client.get("/grades").json()
        assert len(before) == 160
        assert client.post("/imports/grades/commit", json=payload).json()["skipped"] == 160
        assert client.get("/grades").json() == before
        with sessions() as db:
            records = db.scalars(select(GradeRecord)).all()
            assert all(
                record.assessment_config and record.assessment_weight == 1 for record in records
            )


@pytest.mark.parametrize(
    "fill,changes,total",
    [
        ("blank", {}, 0),
        ("blank", {"C7": 0}, 1),
        ("first", {}, 10),
    ],
)
def test_blanks_cached_zero_literal_and_explicit_zero(sessions, fill, changes, total):
    setup(sessions)
    with client_for(sessions) as client:
        result = preview(client, workbook_bytes(fill, changes))
        assert result["total"] == total
        assert result["institutional"]["blank_inputs"] == 160 - total
        assert any(cell["cached"] == "C" for cell in result["institutional"]["derived"])
        if total == 1:
            assert result["rows"][0]["score"] == 0
            assert reference_results(result["rows"])[0]["final"] is None


def test_correct_invalid_preview_then_confirm_keeps_edits(sessions):
    payload = setup(sessions)
    with client_for(sessions) as client:
        result = preview(client, workbook_bytes(changes={"C7": 25}))
        assert result["errors"]
        payload["rows"] = result["rows"]
        assert client.post("/imports/grades/review", json=payload).status_code == 422
        assert client.get("/grades").json() == []
        payload["rows"][0]["score"] = 0
        review = client.post("/imports/grades/review", json=payload)
        assert review.status_code == 200, review.text
        assert not review.json()["errors"]
        assert client.post("/imports/grades/commit", json=payload).json()["imported"] == 10
        assert any(record["score"] == 0 for record in client.get("/grades").json())


def test_student_ambiguous_or_mismatch_requires_explicit_mapping(sessions):
    payload = setup(sessions)
    with client_for(sessions) as client:
        result = preview(
            client, workbook_bytes(changes={"A7": "EXTERNO", "B7": "SINTETICO, AMBIGUO"})
        )
        payload["rows"] = result["rows"]
        assert client.post("/imports/grades/commit", json=payload).status_code == 422
        for row in payload["rows"]:
            row["student_code"] = "1SA001"
        assert client.post("/imports/grades/review", json=payload).status_code == 422
        for row in payload["rows"]:
            row["student_confirmed"] = True
        assert client.post("/imports/grades/review", json=payload).status_code == 200
        assert client.post("/imports/grades/commit", json=payload).json()["imported"] == 10


@pytest.mark.parametrize("change", ["unconfirmed", "year", "date", "order", "section"])
def test_context_required_and_calendar_never_guessed(sessions, change):
    payload = setup(sessions)
    payload["rows"] = read_institutional(workbook_bytes())["rows"]
    ctx = payload["academic_context"]
    if change == "unconfirmed":
        ctx["confirmed"] = False
    elif change == "year":
        ctx["school_year"] = 2025
    elif change == "date":
        ctx["blocks"]["1"]["assessment_date"] = ""
    elif change == "order":
        ctx["blocks"]["2"]["assessment_date"] = "2026-03-01"
    else:
        ctx["section_id"] = ctx["grade_id"]
    with client_for(sessions) as client:
        assert client.post("/imports/grades/commit", json=payload).status_code == 422
        assert client.get("/grades").json() == []


def test_explicit_update_no_silent_replacement_snapshot_and_rollback(sessions):
    payload = setup(sessions)
    payload["rows"] = read_institutional(workbook_bytes())["rows"]
    with client_for(sessions) as client:
        assert client.post("/imports/grades/commit", json=payload).status_code == 200
        before = client.get("/grades").json()
        with sessions() as db:
            course = db.get(Course, payload["course_id"])
            db.add(CourseAssessmentComponent(course_id=course.id, name="Cuaderno", weight=9))
            db.commit()
        updated = deepcopy(payload)
        updated["rows"][0]["score"] = 18
        assert client.post("/imports/grades/commit", json=updated).status_code == 409
        assert client.get("/grades").json() == before
        review = client.post("/imports/grades/review", json=updated).json()
        for row, item in zip(updated["rows"], review["rows"], strict=True):
            row.update(
                component_id=item["component_id"],
                review_token=item["review_token"],
                decision="update",
            )
        assert client.post("/imports/grades/commit", json=updated).json()["updated"] == 1
        with sessions() as db:
            assert all(r.assessment_weight == 1 for r in db.scalars(select(GradeRecord)))
        before_failure = client.get("/grades").json()
        updated["rows"].append(dict(updated["rows"][0], student_code="NO-EXISTE"))
        assert client.post("/imports/grades/commit", json=updated).status_code == 422
        assert client.get("/grades").json() == before_failure


def test_forbidden_course_and_student_role(sessions):
    payload = setup(sessions)
    payload["rows"] = read_institutional(workbook_bytes())["rows"]
    with client_for(sessions, role="teacher") as client:
        assert client.post("/imports/grades/review", json=payload).status_code == 422
        assert client.post("/imports/grades/commit", json=payload).status_code == 422
    with client_for(sessions, role="student") as client:
        assert client.post("/imports/grades/commit", json=payload).status_code == 403


def test_unknown_or_unsafe_formula_not_executed(sessions):
    setup(sessions)
    with client_for(sessions) as client:
        result = preview(
            client, workbook_bytes(changes={"H7": '=WEBSERVICE("https://invalid.example")'})
        )
        assert result["total"] == 10
        assert not result["rows"][0]["reference_verified"]
        assert reference_results(result["rows"])[0]["final"] is None
        result = preview(client, workbook_bytes(changes={"C7": "=1+1"}))
        assert result["errors"]  # Formula in input cannot become a registered grade.


def test_layout_displaced_without_fixed_columns_or_header_row(sessions):
    setup(sessions)
    workbook = load_workbook(BytesIO(workbook_bytes()))
    sheet = workbook.active
    merges = [str(item) for item in sheet.merged_cells.ranges]
    for merged in merges:
        sheet.unmerge_cells(merged)
    sheet.move_range("A1:P22", rows=2, cols=2, translate=True)
    for merged in merges:
        area = CellRange(merged)
        area.shift(row_shift=2, col_shift=2)
        sheet.merge_cells(str(area))
    buffer = BytesIO()
    workbook.save(buffer)
    with client_for(sessions) as client:
        result = preview(client, buffer.getvalue())
        assert result["total"] == 10
        assert result["rows"][0]["source_cell"] == "E9"
        assert all(row["reference_verified"] for row in result["rows"])


def test_assigned_teacher_can_import_and_update(sessions):
    payload = setup(sessions)
    payload["rows"] = read_institutional(workbook_bytes())["rows"]
    from app.models.user import Role, User

    teacher_id = uuid4()
    with sessions() as db:
        role = Role(code="teacher", name="Docente sintético")
        db.add(role)
        db.flush()
        db.add(
            User(
                id=teacher_id,
                role_id=role.id,
                email="teacher@synthetic.example",
                full_name="Docente sintético",
                password_hash="not-a-real-password",
            )
        )
        db.flush()
        assignment = db.scalar(
            select(CourseAssignment).where(CourseAssignment.course_id == payload["course_id"])
        )
        assignment.teacher_id = teacher_id
        db.commit()
    with client_for(sessions, role="teacher", user_id=teacher_id) as client:
        assert client.post("/imports/grades/commit", json=payload).json()["imported"] == 10
        payload["rows"][0]["score"] = 19
        response = client.post("/imports/grades/review", json=payload)
        assert response.status_code == 200, response.text
        item = response.json()["rows"][0]
        payload["rows"][0].update(review_token=item["review_token"], decision="update")
        assert client.post("/imports/grades/commit", json=payload).json()["updated"] == 1


def test_decimal_rounding_and_strict_literal_thresholds(sessions):
    setup(sessions)
    rows = read_institutional(workbook_bytes())["rows"]
    for row in rows:
        row["score"] = 16.45
    result = reference_results(rows)[0]
    assert result["monthly"] == {"1": 16.5, "2": 16.5}
    assert result["literal"] == "AD"
    for row in rows:
        row["score"] = 16.4
    assert reference_results(rows)[0]["literal"] == "A"

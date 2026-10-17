"""Corrections preserve pupil/grade identities in an isolated PostgreSQL schema."""

import os
from copy import deepcopy
from io import BytesIO

import pytest
from institutional_fixture import workbook_bytes
from openpyxl import load_workbook
from sqlalchemy import func, select
from test_grade_import_identity import client_for, context
from test_institutional_excel import preview, setup
from test_roster_persistence import sessions as _sessions

from app.models.academic import GradeRecord
from app.models.student import Student

sessions = _sessions
pytestmark = pytest.mark.skipif(
    not os.environ.get("PAIDEIA_TEST_POSTGRES_URL"), reason="Requires isolated PostgreSQL"
)


def test_roster_template_and_name_correction(sessions):
    context(sessions)
    with sessions() as db:
        pupil = db.scalar(select(Student))
        pupil_id = pupil.id
        payload = {
            "grade_id": str(pupil.grade_id),
            "section_id": str(pupil.section_id),
            "rows": [
                {
                    "student_code": pupil.student_code,
                    "first_name": "CORREGIDO",
                    "last_name": pupil.last_name,
                }
            ],
        }
        count = db.scalar(select(func.count()).select_from(Student))
    with client_for(sessions, "admin") as client:
        preview = client.post(
            "/imports/preview", files={"file": ("notas.xlsx", workbook_bytes("full"))}
        )
        assert preview.status_code == 200
        assert preview.json()["total"] == 16
        assert "full_name" in preview.json()["columns"]
        result = client.post("/imports/students/commit", json=payload)
        assert result.status_code == 200
        assert result.json()["updated"] == 1
        assert result.json()["imported"] == 0
        assert client.post("/imports/students/commit", json=payload).json()["skipped"] == 1
        invalid = deepcopy(payload)
        invalid["rows"].append({**payload["rows"][0], "student_code": "bad", "first_name": ""})
        invalid["rows"][0]["first_name"] = "NO GUARDAR"
        assert client.post("/imports/students/commit", json=invalid).status_code == 422
    with sessions() as db:
        assert db.scalar(select(func.count()).select_from(Student)) == count
        assert db.get(Student, pupil_id).first_name == "CORREGIDO"


def test_correction_only_skips_new_and_updates_reviewed_grade(sessions):
    payload = context(sessions)
    with client_for(sessions) as client:
        initial = {**payload, "rows": payload["rows"][:1]}
        assert client.post("/imports/grades/commit", json=initial).json()["imported"] == 1
        with sessions() as db:
            record_id = db.scalar(select(GradeRecord.id))
        correction = deepcopy(payload)
        correction["corrections_only"] = True
        correction["rows"][0]["score"] = 19
        review = client.post("/imports/grades/review", json=correction).json()
        correction["rows"][0].update(
            decision="update", review_token=review["rows"][0]["review_token"]
        )
        result = client.post("/imports/grades/commit", json=correction)
        assert result.status_code == 200
        assert result.json()["updated"] == 1
        assert result.json()["imported"] == 0
        assert result.json()["skipped"] == 9
    with sessions() as db:
        assert db.scalar(select(func.count()).select_from(GradeRecord)) == 1
        assert db.get(GradeRecord, record_id).score == 19


def test_excel_score_correction_preserves_system_observation(sessions):
    payload = setup(sessions)
    with client_for(sessions) as client:
        payload["rows"] = preview(client, workbook_bytes())["rows"]
        assert client.post("/imports/grades/commit", json=payload).status_code == 200
        with sessions() as db:
            saved = db.scalar(select(GradeRecord).where(GradeRecord.score == 10))
            record_id = saved.id
            saved.qualitative_note = "Observación del sistema"
            db.commit()
        review = client.post("/imports/grades/review", json=payload).json()
        assert all(item["status"] == "identical" for item in review["rows"])
        payload["rows"][0]["score"] = 11
        review = client.post("/imports/grades/review", json=payload).json()
        payload["rows"][0].update(decision="update", review_token=review["rows"][0]["review_token"])
        payload["corrections_only"] = True
        assert client.post("/imports/grades/commit", json=payload).json()["updated"] == 1
    with sessions() as db:
        saved = db.get(GradeRecord, record_id)
        assert saved.score == 11
        assert saved.qualitative_note == "Observación del sistema"


def test_same_surname_first_workbook_imports_pupils_then_grades(sessions):
    payload = setup(sessions)
    workbook = load_workbook(BytesIO(workbook_bytes("full")))
    for source_row in range(7, 23):
        order = source_row + 10
        workbook.active[f"A{source_row}"] = f"1SA{order:03d}"
        workbook.active[f"B{source_row}"] = f"APELLIDO SINTETICO ALUMNO {order:03d}"
    content = BytesIO()
    workbook.save(content)
    with client_for(sessions) as client:
        payload["rows"] = preview(client, content.getvalue())["rows"]
        assert client.post("/imports/grades/review", json=payload).status_code == 422
        with client_for(sessions, "admin") as admin:
            roster = admin.post(
                "/imports/preview", files={"file": ("registro.xlsx", content.getvalue())}
            )
            assert roster.status_code == 200
            rows = roster.json()["rows"]
            assert len(rows) == 16
            assert all(row["last_name"] == "APELLIDO SINTETICO" for row in rows)
            assert all(row["first_name"].startswith("ALUMNO ") for row in rows)
            assert roster.json()["name_split_warning"]
            student_payload = {**payload["academic_context"], "rows": rows}
            result = admin.post("/imports/students/commit", json=student_payload)
            assert result.status_code == 200, result.text
            assert result.json()["imported"] == 16
            assert admin.post("/imports/students/commit", json=student_payload).json()["skipped"] == 16
        reviewed = client.post("/imports/grades/review", json=payload)
        assert reviewed.status_code == 200, reviewed.text
        assert reviewed.json()["total"] == 160
        assert not reviewed.json()["errors"]
        assert (
            client.post(
                "/imports/grades/commit", json={**payload, "corrections_only": True}
            ).json()["imported"]
            == 0
        )
        assert client.post("/imports/grades/commit", json=payload).json()["imported"] == 160
        assert client.post("/imports/grades/commit", json=payload).json()["skipped"] == 160


def test_existing_compound_surname_boundary_is_preserved(sessions):
    setup(sessions)
    with sessions() as db:
        student = db.scalar(select(Student).where(Student.student_code == "1SA001"))
        student.last_name = "DE LA CRUZ SINTETICO"
        student.first_name = "ALUMNO UNO"
        db.commit()
    workbook = load_workbook(BytesIO(workbook_bytes()))
    workbook.active["B7"] = "DE LA CRUZ SINTETICO ALUMNO UNO"
    content = BytesIO()
    workbook.save(content)
    with client_for(sessions, "admin") as client:
        result = client.post(
            "/imports/preview", files={"file": ("registro.xlsx", content.getvalue())}
        )
        assert result.status_code == 200
        assert result.json()["rows"][0]["last_name"] == "DE LA CRUZ SINTETICO"
        assert result.json()["rows"][0]["first_name"] == "ALUMNO UNO"

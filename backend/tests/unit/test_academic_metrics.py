"""Missing evidence, captured weights, genuine zero and consistent scoped exports."""

from io import BytesIO
from uuid import UUID, uuid4

import pytest
from openpyxl import load_workbook
from test_academic_authorization import scenario, sessions

from app.models.academic import CourseAssessmentComponent
from app.models.student import Student

__all__ = ["scenario", "sessions"]


@pytest.mark.parametrize("score,risk", [(0, "ALTO"), (18, "INCOMPLETO")])
def test_valid_scores_without_attendance_keep_their_value(scenario, score, risk):
    client, ids, _actors, _current = scenario
    response = client.post(
        "/grades",
        json={
            "student_id": ids["students"][0],
            "course_id": ids["courses"][0],
            "period_id": ids["period"],
            "evaluation_name": "Solo notas",
            "assessment_date": "2026-05-01",
            "score": score,
        },
    )
    assert response.status_code == 201
    row = client.get(
        "/reports/students", params={"month": 5, "course_id": ids["courses"][0]}
    ).json()[0]
    assert row["average_score"] == score
    assert row["attendance_rate"] is None
    assert row["risk_level"] == risk


def test_empty_filtered_reports_are_not_zero_or_high_risk(scenario):
    client, ids, _actors, _current = scenario
    rows = client.get("/reports/students", params={"month": 2}).json()
    assert len(rows) == 2
    assert all(row["average_score"] is None and row["attendance_rate"] is None for row in rows)
    assert all(row["risk_level"] == "SIN_DATOS" for row in rows)
    assert all(row["grade_count"] == row["attendance_count"] == 0 for row in rows)
    summary = client.get("/dashboard/summary", params={"month": 2}).json()
    assert summary["average_score"] is summary["attendance_rate"] is None
    assert summary["at_risk_count"] == 0
    result = client.get("/reports/export.xlsx", params={"grade_id": ids["grade"], "month": 2})
    sheet = load_workbook(BytesIO(result.content)).active
    columns = {cell.value: cell.column for cell in sheet[1]}
    assert sheet.cell(2, columns["average_score"]).value is None
    assert sheet.cell(2, columns["attendance_rate"]).value is None
    assert sheet.cell(2, columns["risk_level"]).value == "SIN_DATOS"


def test_zero_and_snapshot_weights_are_shared_by_reports_dashboard_and_summary(scenario, sessions):
    client, ids, _actors, _current = scenario
    with sessions() as db:
        component = CourseAssessmentComponent(
            id=uuid4(), course_id=UUID(ids["courses"][0]), name="Peso histórico", weight=3
        )
        db.add(component)
        db.commit()
        component_id = component.id
    created = client.post(
        "/grades",
        json={
            "student_id": ids["students"][0],
            "course_id": ids["courses"][0],
            "period_id": ids["period"],
            "component_id": str(component_id),
            "evaluation_name": "Cero válido",
            "assessment_date": "2026-04-02",
            "score": 0,
        },
    )
    assert created.status_code == 201, created.text
    assert created.json()["assessment_weight"] == 3
    with sessions() as db:
        db.get(CourseAssessmentComponent, component_id).weight = 9
        db.commit()
    rows = client.get("/reports/students").json()
    first = next(row for row in rows if row["student_id"] == ids["students"][0])
    assert first["average_score"] == 2.5
    assert first["grade_weight_sum"] == 4 and first["grade_count"] == 2
    assert client.get(f"/grades/summary/{ids['students'][0]}").json()["average"] == 2.5
    for route in ("/dashboard/summary", "/reports/summary"):
        summary = client.get(route).json()
        assert summary["average_score"] == 4.4
        assert summary["attendance_rate"] == 100
        assert summary["at_risk_count"] == 1
    zero_only = client.get(
        "/reports/students", params={"month": 4, "course_id": ids["courses"][0]}
    ).json()[0]
    assert zero_only["risk_level"] == "ALTO"
    assert (
        client.get("/grades").json()[0]["assessment_config"] == created.json()["assessment_config"]
    )


def test_student_without_any_records_has_null_average_and_is_not_an_alert(scenario, sessions):
    client, ids, _actors, _current = scenario
    with sessions() as db:
        student = Student(
            id=uuid4(),
            student_code="SYN-NODATA",
            first_name="Sin",
            last_name="Evidencia",
            grade_id=UUID(ids["grade"]),
            section_id=UUID(ids["sections"][0]),
        )
        db.add(student)
        db.commit()
        student_id = str(student.id)
    assert client.get(f"/grades/summary/{student_id}").json() == {
        "student_id": student_id,
        "average": None,
        "count": 0,
    }
    row = next(r for r in client.get("/reports/students").json() if r["student_id"] == student_id)
    assert row["risk_level"] == "SIN_DATOS"
    assert (
        client.get("/reports/summary").json()["at_risk_count"] == 1
    )  # existing 10, not missing student


def test_zero_weight_has_no_defined_average_and_does_not_fabricate_zero(scenario, sessions):
    client, ids, _actors, _current = scenario
    with sessions() as db:
        component = CourseAssessmentComponent(
            id=uuid4(), course_id=UUID(ids["courses"][0]), name="Sin ponderación", weight=0
        )
        db.add(component)
        db.commit()
        component_id = str(component.id)
    response = client.post(
        "/grades",
        json={
            "student_id": ids["students"][0],
            "course_id": ids["courses"][0],
            "period_id": ids["period"],
            "component_id": component_id,
            "evaluation_name": "Peso cero",
            "assessment_date": "2026-05-01",
            "score": 0,
        },
    )
    assert response.status_code == 201
    row = client.get(
        "/reports/students", params={"month": 5, "course_id": ids["courses"][0]}
    ).json()[0]
    assert row["average_score"] is None and row["grade_count"] == 1 and row["grade_weight_sum"] == 0
    monthly = client.get(
        "/grades/monthly-summary",
        params={
            "grade_id": ids["grade"],
            "section_id": ids["sections"][0],
            "course_id": ids["courses"][0],
            "period_id": ids["period"],
            "month": 5,
        },
    ).json()["students"][0]
    assert monthly["average"] is monthly["literal"] is None

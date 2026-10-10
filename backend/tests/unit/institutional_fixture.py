"""Synthetic replica of the inspected anonymous Hoja1 (22x16, nine merges)."""

from io import BytesIO
from xml.etree import ElementTree as ET
from zipfile import ZipFile

from openpyxl import Workbook


def workbook_bytes(fill="first", changes=None):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Hoja1"
    for merged in ["A1:O2", "A3:B4", "C3:O4", "A5:A6", "B5:B6", "C5:H5", "I5:N5", "O5:O6", "P5:P6"]:
        sheet.merge_cells(merged)
    for position, value in {
        "A1": "REGISTRO DE NOTAS ",
        "A3": "SINTETICO",
        "C3": "SINTETICO",
        "A5": "CODIGO",
        "B5": "NOMBRE Y APELLIDOS",
        "C5": "1 MES",
        "I5": "2 MES",
        "O5": "P.F",
        "P5": "PROMEDIO LITERAL",
    }.items():
        sheet[position] = value
    criteria = ["ACTITUD ANTE EL AREA", "CUADERNO", "MODULO", "EXPOSICION-TRABAJOS", "EVALUACION"]
    for start in (3, 9):
        for offset, header in enumerate(criteria):
            sheet.cell(6, start + offset, header)
    sheet["H6"], sheet["N6"] = "PROMEDIO MENSUAL", "P.S"
    for row in range(7, 23):
        sheet[f"A{row}"] = f"1SA{row - 6:03d}"
        sheet[f"B{row}"] = f"SINTETICO, ALUMNO {row - 6:03d}"
        if fill == "full" or (fill == "first" and row == 7):
            for column, value in zip(
                "CDEFGIJKLM", [10, 12, 14, 16, 18, 20, 20, 20, 20, 20], strict=True
            ):
                sheet[f"{column}{row}"] = value
        sheet[f"H{row}"] = f"=ROUND((SUM(C{row}:E{row})+F{row}*1+G{row}*1)/5,1)"
        sheet[f"N{row}"] = f"=ROUND(((SUM(I{row}:K{row}))+L{row}*1+M{row}*1)/5,1)"
        sheet[f"O{row}"] = f"=ROUND(AVERAGE(H{row},N{row}),1)"
        sheet[f"P{row}"] = f'=IF(O{row}>16.4,"AD",IF(O{row}>13.4,"A",IF(O{row}>10.4,"B","C")))'
    for position, value in (changes or {}).items():
        sheet[position] = value
    output = BytesIO()
    workbook.save(output)
    # Excel caches from blank inputs may contain 0/C: faithfully reproduce that trap.
    result = BytesIO()
    namespace = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
    with ZipFile(output) as source, ZipFile(result, "w") as destination:
        for item in source.infolist():
            content = source.read(item.filename)
            if item.filename == "xl/worksheets/sheet1.xml":
                root = ET.fromstring(content)
                for cell in root.iter(f"{{{namespace}}}c"):
                    if cell.find(f"{{{namespace}}}f") is None:
                        continue
                    cached = cell.find(f"{{{namespace}}}v")
                    if cached is None:
                        cached = ET.SubElement(cell, f"{{{namespace}}}v")
                    if cell.get("r").startswith("P"):
                        cell.set("t", "str")
                        cached.text = "C"
                    else:
                        cached.text = "0"
                content = ET.tostring(root)
            destination.writestr(item, content)
    return result.getvalue()
